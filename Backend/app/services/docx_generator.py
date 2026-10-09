from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor

from app.schemas.resume import ResumeData


def build_docx(r: ResumeData) -> bytes:
    doc = Document()
    doc.add_heading(r.name or "Your Name", level=0)
    info = " | ".join(x for x in [r.email, r.phone, r.location] if x)
    if info:
        doc.add_paragraph(info)

    def heading(text):
        doc.add_heading(text, level=1)

    def entry(title, meta, desc):
        p = doc.add_paragraph()
        run = p.add_run(title)
        run.bold = True
        if meta:
            m = p.add_run(f"\n{meta}")
            m.font.size = Pt(9)
            m.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        if desc:
            p.add_run(f"\n{desc}")

    if r.summary:
        heading("Summary")
        doc.add_paragraph(r.summary)
    if r.skills:
        heading("Skills")
        doc.add_paragraph(", ".join(r.skills))
    if r.experience:
        heading("Experience")
        for x in r.experience:
            entry(f"{x.role} — {x.company}", " - ".join(d for d in [x.start_date, x.end_date] if d), x.description)
    if r.projects:
        heading("Projects")
        for p in r.projects:
            meta = ("Tech: " + ", ".join(p.technologies)) if p.technologies else ""
            entry(p.name, meta, "\n".join(x for x in [p.description, p.link] if x))
    if r.education:
        heading("Education")
        for x in r.education:
            entry(f"{x.degree} — {x.institution}", " - ".join(d for d in [x.start_year, x.end_year] if d), "")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
